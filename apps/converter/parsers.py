"""Fayl parserlari — Word / Excel / TXT → savollar ro'yxati.

Kutilayotgan format (har qanday faylda):

1. Savol matni
A) Variant 1
*B) Variant 2 (yulduzcha — to'g'ri javob)
C) Variant 3
D) Variant 4

2. Keyingi savol matni
...
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Iterable

from .normalizer import normalize_for_hemis


@dataclass
class ParsedChoice:
    text: str
    is_correct: bool = False


@dataclass
class ParsedQuestion:
    text: str
    choices: list[ParsedChoice] = field(default_factory=list)

    @property
    def is_valid(self) -> bool:
        return bool(self.text) and len(self.choices) >= 2 and any(c.is_correct for c in self.choices)


Q_NUM = re.compile(r'^\s*(\d+)[.)]\s*(.*)$')
# To'g'ri javob belgilari: * ** + ++ # ## (variant harfidan oldin yoki keyin)
CHOICE_RE = re.compile(
    r'^\s*(?P<pre>[*+#]{1,2})?\s*(?P<letter>[A-Za-z])\s*(?P<post>[*+#]{1,2})?\s*[.)]\s*(?P<text>.*)$'
)
ANSWER_RE = re.compile(r'^\s*(?:ANSWER|JAVOB)\s*[:=]\s*([A-Za-z,\s]+)\s*$', re.I)
LEADING_MARK = re.compile(r'^([*+#]{1,2})\s*(.+)$')


def _expand_inline_choices(line: str) -> list[str]:
    """Bir qatorda bir nechta variant bo'lsa (A) .. B) .. C) ..) — bo'lib chiqaradi.

    A, B, C, D ketma-ketligini topadi — variantlar oralig'ida probel
    bo'lmasa ham (Word soft line-break holati) ishlaydi.
    """
    candidates = list(re.finditer(r'([*+#]{0,2})\s*([A-Za-z])\s*([*+#]{0,2})[.)]', line))
    if len(candidates) < 2:
        return [line]
    # A dan boshlanadigan ketma-ket ABC.. zanjirini topish
    positions: list[int] = []
    expected = ord('A')
    for m in candidates:
        L = m.group(2).upper()
        if not positions:
            if L == 'A':
                positions.append(m.start())
                expected = ord('B')
            continue
        if ord(L) == expected:
            positions.append(m.start())
            expected += 1
    if len(positions) < 2:
        return [line]
    pieces: list[str] = []
    prefix = line[:positions[0]].strip()
    if prefix:
        pieces.append(prefix)
    for idx, p in enumerate(positions):
        end = positions[idx + 1] if idx + 1 < len(positions) else len(line)
        piece = line[p:end].strip()
        if piece:
            pieces.append(piece)
    return pieces


def _parse_lines(lines: Iterable[str]) -> list[ParsedQuestion]:
    questions: list[ParsedQuestion] = []
    current: ParsedQuestion | None = None
    letter_to_choice: dict[str, ParsedChoice] = {}

    def flush():
        nonlocal current, letter_to_choice
        if current and current.text:
            questions.append(current)
        current = None
        letter_to_choice = {}

    expanded: list[str] = []
    for raw in lines:
        rs = raw.rstrip()
        bold_prefix = '\x01' if rs.startswith('\x01') else ''
        body = rs[1:] if bold_prefix else rs
        for piece in _expand_inline_choices(body):
            expanded.append(bold_prefix + piece if bold_prefix else piece)

    for raw in expanded:
        line = raw.rstrip()
        bold_correct = line.startswith('\x01')
        if bold_correct:
            line = line[1:]
        if not line.strip():
            continue

        m_q = Q_NUM.match(line)
        m_c = CHOICE_RE.match(line)
        m_a = ANSWER_RE.match(line)

        if m_a and current:
            letters = [x.strip().upper() for x in m_a.group(1).split(',') if x.strip()]
            for L in letters:
                if L in letter_to_choice:
                    letter_to_choice[L].is_correct = True
            # Aiken format: ANSWER qatori — savol tugadi, keyingisini boshlaymiz
            flush()
            continue

        # Variant emas, savol raqami emas, ammo joriy savol allaqachon
        # variantlarga ega — bu yangi savolning matni (Aiken: raqamsiz)
        if not m_c and not m_q and current and current.choices:
            flush()
            current = ParsedQuestion(text=line.strip())
            continue

        if m_q and not m_c:
            flush()
            current = ParsedQuestion(text=m_q.group(2).strip())
            continue

        if m_c:
            if current is None:
                current = ParsedQuestion(text="")
            letter = m_c.group('letter')
            text = m_c.group('text')
            is_correct = bool(m_c.group('pre') or m_c.group('post')) or bold_correct
            # Variant matni "# ..." / "* ..." / "+ ..." bilan boshlansa — to'g'ri javob
            m_lead = LEADING_MARK.match(text.strip())
            if m_lead:
                is_correct = True
                text = m_lead.group(2)
            choice = ParsedChoice(text=text.strip(), is_correct=is_correct)
            current.choices.append(choice)
            letter_to_choice[letter.upper()] = choice
            continue

        # Oddiy matn — joriy savol matniga qo'shish
        if current is None:
            current = ParsedQuestion(text=line.strip())
        elif not current.choices:
            current.text = (current.text + " " + line.strip()).strip()

    flush()

    # HEMIS uchun har bir matnni normalizatsiya qilish
    for q in questions:
        q.text = normalize_for_hemis(q.text)
        for c in q.choices:
            c.text = normalize_for_hemis(c.text)
    return questions


def parse_text(content: str) -> list[ParsedQuestion]:
    return _parse_lines(content.splitlines())


def parse_docx(file) -> list[ParsedQuestion]:
    try:
        from docx import Document
    except ImportError as e:
        raise RuntimeError(
            "python-docx o'rnatilmagan. Terminalda: pip install python-docx, "
            "so'ng serverni qayta ishga tushiring."
        ) from e
    doc = Document(file)

    def para_lines(p):
        text = p.text
        if not text.strip():
            return ['']
        # Paragraf ichida soft line-break (\n) bo'lsa — alohida qatorlarga ajratish
        sub = text.splitlines()
        # Birinchi non-empty run bold bo'lsa — birinchi qatorga sentinel
        bold = False
        for run in p.runs:
            if not run.text.strip():
                continue
            bold = bool(run.bold)
            break
        if bold and sub:
            sub[0] = '\x01' + sub[0]
        return sub

    lines: list[str] = []
    for p in doc.paragraphs:
        lines.extend(para_lines(p))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.text.strip():
                        lines.extend(para_lines(p))
    return _parse_lines(lines)


def parse_xlsx(file) -> list[ParsedQuestion]:
    """Excel format: A=savol, B=to'g'ri_javob, C..F=variantlar."""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError("openpyxl o'rnatilmagan. pip install openpyxl") from e
    wb = load_workbook(file, data_only=True, read_only=True)
    ws = wb.active
    questions: list[ParsedQuestion] = []
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    # Agar birinchi qatorda 'savol' yoki 'question' bo'lsa — sarlavha deb o'tkazib yuborish
    first = [str(c or '').lower() for c in rows[0]]
    start = 1 if any(k in ' '.join(first) for k in ('savol', 'question')) else 0

    for row in rows[start:]:
        if not row or not row[0]:
            continue
        q_text = str(row[0]).strip()
        correct = str(row[1] or '').strip().upper() if len(row) > 1 else 'A'
        choices: list[ParsedChoice] = []
        for idx, val in enumerate(row[2:6]):
            if val is None or str(val).strip() == '':
                continue
            letter = 'ABCDEFGH'[idx]
            choices.append(ParsedChoice(
                text=str(val).strip(),
                is_correct=(letter in correct),
            ))
        if q_text and choices:
            questions.append(ParsedQuestion(text=q_text, choices=choices))
    return questions


def _words_to_lines(words: list) -> list[str]:
    """pdfplumber extract_words natijasidan y-koordinata bo'yicha qatorlarni qayta qurish."""
    Y_CLUSTER = 3  # bir xil qatordagi so'zlar orasidagi max top-farqi (pt)
    clusters: list[tuple[float, list]] = []
    for word in sorted(words, key=lambda w: w['top']):
        merged = False
        for cluster in reversed(clusters):
            if abs(word['top'] - cluster[0]) <= Y_CLUSTER:
                cluster[1].append(word)
                merged = True
                break
        if not merged:
            clusters.append((word['top'], [word]))
    lines: list[str] = []
    for _, cluster_words in sorted(clusters, key=lambda c: c[0]):
        lines.append(' '.join(w['text'] for w in sorted(cluster_words, key=lambda w: w['x0'])))
    return lines


def _extract_pdf_page_lines(page) -> list[str]:
    # 1. layout=True — pdfminer layout engine, ko'p ustunli va murakkab PDFlar uchun
    try:
        text = page.extract_text(layout=True, x_tolerance=3, y_tolerance=3) or ''
    except TypeError:
        # eski pdfplumber versiyasida layout parametri yo'q
        text = page.extract_text(x_tolerance=3, y_tolerance=3) or ''
    if text.strip():
        return text.splitlines()

    # 2. So'zma-so'z qayta qurish — koordinatalar bo'yicha to'g'ri tartib
    words = page.extract_words(x_tolerance=3, y_tolerance=3, keep_blank_chars=False)
    if words:
        return _words_to_lines(words)

    # 3. Jadval sifatida o'qib ko'rish
    lines: list[str] = []
    for table in (page.extract_tables() or []):
        for row in table:
            cells = [str(c or '').strip() for c in row if c and str(c).strip()]
            if cells:
                lines.append(' '.join(cells))
    return lines


def parse_pdf(file) -> list[ParsedQuestion]:
    try:
        import pdfplumber
    except ImportError as e:
        raise RuntimeError(
            "pdfplumber o'rnatilmagan. Terminalda: pip install pdfplumber, "
            "so'ng serverni qayta ishga tushiring."
        ) from e

    all_lines: list[str] = []
    with pdfplumber.open(file) as pdf:
        for page in pdf.pages:
            all_lines.extend(_extract_pdf_page_lines(page))

    if not any(ln.strip() for ln in all_lines):
        raise RuntimeError(
            "PDF dan matn o'qib bo'lmadi. "
            "Fayl skanerlangan rasm bo'lishi mumkin — "
            "faylni DOCX yoki TXT formatiga o'girib qayta yuklang."
        )

    return _parse_lines(all_lines)


def parse_file(file, filename: str) -> list[ParsedQuestion]:
    name = filename.lower()
    if name.endswith('.docx'):
        return parse_docx(file)
    if name.endswith('.doc'):
        raise ValueError(
            "Eski Word formati (.doc) to'g'ridan-to'g'ri o'qib bo'lmaydi. "
            "Faylni Microsoft Word'da ochib: 'Fayl' → 'Boshqacha saqlash' → '.docx' formatini tanlang. "
            "Yoki AI konvertori yoqilgan bo'lsa avtomatik ishlab chiqiladi."
        )
    if name.endswith('.xlsx'):
        return parse_xlsx(file)
    if name.endswith('.xls'):
        raise ValueError(
            "Eski Excel formati (.xls) qo'llab-quvvatlanmaydi. "
            "Faylni Excel'da ochib: 'Fayl' → 'Boshqacha saqlash' → '.xlsx' formatini tanlang."
        )
    if name.endswith('.pdf'):
        return parse_pdf(file)
    if name.endswith('.txt'):
        content = file.read()
        if isinstance(content, bytes):
            content = content.decode('utf-8', errors='replace')
        return parse_text(content)
    raise ValueError(
        f"'{filename}' formati qo'llab-quvvatlanmaydi. "
        "Qabul qilinadigan formatlar: .docx, .xlsx, .pdf, .txt"
    )
