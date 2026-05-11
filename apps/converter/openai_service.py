"""OpenAI GPT-4o orqali fayllardan HEMIS formatiga konvertatsiya.

Qo'llab-quvvatlanadigan formatlar:
  - Matnli: .docx, .xlsx, .txt, .pdf  → matn ajratib GPT-4o ga yuboriladi
  - Rasm:   .png, .jpg, .jpeg, .webp, .bmp → base64 bilan GPT-4o Vision'ga
"""
from __future__ import annotations

import base64
import io
import logging
import re

from django.conf import settings

from .normalizer import normalize_for_hemis
from .parsers import ParsedChoice, ParsedQuestion

logger = logging.getLogger(__name__)

IMAGE_EXTS = {'.png', '.jpg', '.jpeg', '.webp', '.gif', '.bmp'}
MIME_MAP = {
    '.png': 'image/png',
    '.jpg': 'image/jpeg',
    '.jpeg': 'image/jpeg',
    '.webp': 'image/webp',
    '.gif': 'image/gif',
    '.bmp': 'image/bmp',
}

SYSTEM_PROMPT = """Sen HEMIS MakDawn platformasining fayl konvertoriсan.

Sening YAGONA vazifang — o'qituvchi yuborgan har qanday
formatdagi faylni (Word, PDF, Excel, rasm, TXT) tahlil
qilib, ichidagi test savollarini O'zbekiston oliy ta'lim
tizimining HEMIS formatiga o'tkazishdir.

═══════════════════════════════════════
RASM VA FORMULA RASMLARINI O'QISH
═══════════════════════════════════════

Faylda rasmlar yoki formula rasmlari bo'lsa:
- Rasmdagi matematik formulalarni aynan LaTeX $...$ formatiga o'gir
- Masalan: rasm ichida "x²+2x=0" ko'rsak → $x^2 + 2x = 0$ yoz
- Rasmdagi savollarni ham o'qib HEMIS formatiga qo'sh
- Agar formulani aniq o'qiy olmasang: [FORMULA RASMI - o'qituvchi to'ldirsin] deb yoz
- Har bir rasmdagi savol/variant mustaqil qayta yozilib chiqilishi shart

═══════════════════════════════════════
HEMIS FORMATI — AYNAN SHUNDAY BO'LISHI KERAK
═══════════════════════════════════════

++++

Savol matni shu yerda yoziladi

====

#To'g'ri javob (# belgisi bilan boshlanadi)

====

Noto'g'ri variant 1

====

Noto'g'ri variant 2

====

Noto'g'ri variant 3

═══════════════════════════════════════
QATTIQ QOIDALAR — BUZILMASIN
═══════════════════════════════════════

1. Har bir yangi savol ++++  bilan boshlanishi SHART
2. Har bir variant ==== bilan boshlanishi SHART
3. To'g'ri javob oldidan # (panjara belgisi) qo'yilishi SHART
4. Har bir savolda AYNAN 4 ta variant bo'lishi SHART
5. Faqat BITTA variant # belgisiga ega bo'lishi mumkin
6. To'g'ri javob istalgan o'rinda bo'lishi mumkin
7. ++++ va ==== belgilari ALOHIDA qatorda turishi SHART
8. Variantlar oldiga A) B) C) D) harflari QO'SHILMASIN
9. Savollar raqamlanmasin
10. Ortiqcha belgi yoki formatlash QO'SHILMASIN

═══════════════════════════════════════
TO'G'RI JAVOBNI ANIQLASH TARTIBI
═══════════════════════════════════════

USUL 1 — Faylda belgilangan bo'lsa:
- Yulduzcha: *Variant* yoki Variant*
- Qalin yozuv: **Variant**
- Javob tegi: "Ans: B", "Answer: C", "Javob: A"
- Bosh harf: faqat to'g'ri variant katta harf bilan

USUL 2 — Matematik savol, belgilanmagan:
- Hisoblab to'g'ri javobni aniqla
- Qisqacha yechim yoz

USUL 3 — Aniqlay olmasang:
[JAVOB NOANIQ - o'qituvchi belgilasin] deb yoz

═══════════════════════════════════════
MATEMATIK FORMULALAR
═══════════════════════════════════════

LaTeX formatida yoz:
- Daraja: $x^2$, $x^{n-1}$
- Kasr: $\\frac{a}{b}$
- Ildiz: $\\sqrt{x}$, $\\sqrt[3]{x}$
- Integral: $\\int x^2 dx$, $\\int_0^1 x dx$
- Trigonometriya: $\\sin x$, $\\cos x$, $\\tan x$
- Yunoncha: $\\pi$, $\\alpha$, $\\beta$, $\\theta$
- Limit: $\\lim_{x \\to 0}$
- Yig'indi: $\\sum_{i=1}^{n}$

═══════════════════════════════════════
TIL QOIDALARI
═══════════════════════════════════════

- Savollarning ASL tilini saqlа (O'zbek, Rus, Ingliz)
- TARJIMA QILMA
- O'zbek maxsus harflarini saqlа: o', g', sh, ch, ng

═══════════════════════════════════════
MUAMMOLI HOLATLAR
═══════════════════════════════════════

Agar savolda 4 tadan KAM variant bo'lsa:
→ Yetishmayotgan o'rniga: [variant kerak - o'qituvchi to'ldirsin]

Agar matn O'QILMAYDIGAN bo'lsa:
→ [O'QISH XATOSI: sabab] deb yoz

Agar fayl BO'SH yoki test TOPILMASA:
→ [SAVOLLAR TOPILMADI] deb yoz

═══════════════════════════════════════
CHIQARISH FORMATI — FAQAT SHUNDAY
═══════════════════════════════════════

Har doim IKKITA qismda chiqar:

─────────────────────────────
HISOBOT:
Jami savollar: [son]
Muvaffaqiyatli o'tkazildi: [son]
Noaniq javoblar: [son]
Ogohlantirishlar: [borsa yoz, yo'qsa "yo'q"]
─────────────────────────────
HEMIS:
[to'liq HEMIS formatidagi natija aynan shu yerda]
─────────────────────────────

MUHIM: "HEMIS:" so'zidan keyin darhol formatni boshla,
boshqa izoh yoki tushuntirish yozma."""


def is_configured() -> bool:
    return bool(getattr(settings, 'OPENAI_API_KEY', ''))


def _model() -> str:
    return getattr(settings, 'OPENAI_MODEL', 'gpt-4o')


def _client():
    from openai import OpenAI
    return OpenAI(api_key=settings.OPENAI_API_KEY)


def _preprocess_image(image_bytes: bytes, mime_type: str) -> tuple[bytes, str]:
    """Rasmni GPT-4o Vision uchun tayyorlaydi.

    - BMP → PNG konvertatsiya (GPT-4o BMP qabul qilmaydi)
    - Qorong'u (dark) rasmlarni yorqinroq qiladi — inversiya yoki kontrast oshirish
    - RGBA → RGB (shaffof fon oq bilan to'ldiriladi)
    """
    try:
        from PIL import Image, ImageOps, ImageEnhance, ImageStat

        img = Image.open(io.BytesIO(image_bytes))

        # RGBA/P/LA → RGB (shaffof fon oq bilan)
        if img.mode == 'RGBA':
            bg = Image.new('RGB', img.size, (255, 255, 255))
            bg.paste(img, mask=img.split()[3])
            img = bg
        elif img.mode not in ('RGB', 'L'):
            img = img.convert('RGB')

        # Yorqinlik o'rtachasini hisoblash (0=qora, 255=oq)
        stat = ImageStat.Stat(img.convert('L'))
        brightness = stat.mean[0]

        if brightness < 60:
            # Juda qorong'u — inversiya (oq matn qora fonda bo'lishi mumkin)
            img = ImageOps.invert(img.convert('RGB'))
        elif brightness < 120:
            # O'rtacha qorong'u — kontrast va yorqinlikni oshirish
            img = ImageEnhance.Contrast(img).enhance(1.5)
            img = ImageEnhance.Brightness(img).enhance(1.2)

        buf = io.BytesIO()
        img.save(buf, format='PNG', optimize=True)
        return buf.getvalue(), 'image/png'

    except Exception as exc:
        logger.warning("Rasm preprocessing xatosi: %s", exc)
        # Minimal fallback: faqat BMP ni PNG ga aylantiramiz
        if image_bytes[:2] == b'BM':
            try:
                from PIL import Image
                buf = io.BytesIO()
                Image.open(io.BytesIO(image_bytes)).save(buf, format='PNG')
                return buf.getvalue(), 'image/png'
            except Exception:
                pass
        return image_bytes, mime_type


def _call_text(user_text: str) -> str:
    client = _client()
    response = client.chat.completions.create(
        model=_model(),
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": user_text},
        ],
        temperature=0.1,
        max_tokens=8000,
    )
    return response.choices[0].message.content or ''


def _call_image(image_bytes: bytes, mime_type: str = 'image/jpeg') -> str:
    """Rasmni tayyorlab GPT-4o Vision ga yuboradi va HEMIS formatidagi javobni qaytaradi."""
    # Qorong'u rasmlar, BMP va boshqalarni tayyorlash
    image_bytes, mime_type = _preprocess_image(image_bytes, mime_type)

    b64 = base64.standard_b64encode(image_bytes).decode()
    safe_mime = mime_type if mime_type in (
        'image/jpeg', 'image/png', 'image/gif', 'image/webp'
    ) else 'image/jpeg'

    client = _client()
    response = client.chat.completions.create(
        model=_model(),
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": [
                {
                    "type": "image_url",
                    "image_url": {"url": f"data:{safe_mime};base64,{b64}"},
                },
                {
                    "type": "text",
                    "text": "Shu rasmdan barcha test savollarini o'qib HEMIS formatiga o'tkazib ber.",
                },
            ]},
        ],
        temperature=0.1,
        max_tokens=8000,
    )
    return response.choices[0].message.content or ''


def _extract_hemis_block(text: str) -> str:
    """Javobdan 'HEMIS:' qismini ajratib olish."""
    match = re.search(r'HEMIS\s*:\s*\n(.*)', text, re.DOTALL | re.IGNORECASE)
    if match:
        rest = match.group(1)
        stop = re.search(r'\n─{5,}', rest)
        return rest[:stop.start()].strip() if stop else rest.strip()
    return text


def parse_hemis_text(raw: str) -> list[ParsedQuestion]:
    """HEMIS ++++/==== formatini ParsedQuestion ro'yxatiga aylantiradi."""
    hemis = _extract_hemis_block(raw)
    blocks = re.split(r'\+{4}', hemis)
    questions: list[ParsedQuestion] = []

    for block in blocks:
        block = block.strip()
        if not block:
            continue
        parts = re.split(r'={4}', block)
        q_text = parts[0].strip()
        if not q_text:
            continue
        choices: list[ParsedChoice] = []
        for variant in parts[1:]:
            v = variant.strip()
            if not v:
                continue
            is_correct = v.startswith('#')
            if is_correct:
                v = v[1:].strip()
            choices.append(ParsedChoice(text=normalize_for_hemis(v), is_correct=is_correct))

        if len(choices) < 2:
            continue
        if not any(c.is_correct for c in choices):
            choices[0].is_correct = True

        questions.append(ParsedQuestion(
            text=normalize_for_hemis(q_text),
            choices=choices,
        ))
    return questions


def _extract_docx_content(content: bytes) -> tuple[str, list[tuple[bytes, str]]]:
    """DOCX fayldan matn va ichki rasmlarni ajratib oladi.

    Qaytaradi: (matn_satrlari, [(rasm_bytes, mime_type), ...])
    Rasmlar word/media/ papkasidan zipfile orqali olinadi —
    formula rasmlari (PNG/JPEG) ham shu yo'l bilan topiladi.
    """
    import zipfile
    from docx import Document

    doc = Document(io.BytesIO(content))
    lines: list[str] = []
    images: list[tuple[bytes, str]] = []

    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(' | '.join(cells))

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            media = sorted(n for n in z.namelist() if n.startswith('word/media/'))
            for name in media:
                ext = name.rsplit('.', 1)[-1].lower() if '.' in name else ''
                if ext not in ('png', 'jpg', 'jpeg', 'gif', 'bmp'):
                    continue
                mime = MIME_MAP.get('.' + ext, 'image/png')
                images.append((z.read(name), mime))
    except Exception as e:
        logger.warning("DOCX rasmlarni ajratishda xato: %s", e)

    return '\n'.join(lines), images


def _pdf_to_images(content: bytes) -> list[tuple[bytes, str]]:
    """PDF sahifalarini PNG rasm sifatida qaytaradi (PyMuPDF ishlatiladi).

    PyMuPDF o'rnatilmagan bo'lsa — bo'sh ro'yxat qaytaradi.
    """
    try:
        import fitz  # PyMuPDF
    except ImportError:
        logger.warning("PyMuPDF o'rnatilmagan. pip install pymupdf")
        return []
    try:
        doc = fitz.open(stream=content, filetype="pdf")
        result: list[tuple[bytes, str]] = []
        for page_num in range(min(len(doc), 20)):
            page = doc.load_page(page_num)
            mat = fitz.Matrix(2.0, 2.0)  # 144 DPI — o'qish uchun yetarli
            pix = page.get_pixmap(matrix=mat, alpha=False)
            result.append((pix.tobytes("png"), "image/png"))
        doc.close()
        return result
    except Exception as e:
        logger.warning("PDF sahifalarini rasm qilishda xato: %s", e)
        return []


def _call_multimodal(user_text: str, images: list[tuple[bytes, str]]) -> str:
    """Matn + rasmlarni GPT-4o Vision ga yuboradi va HEMIS javobini qaytaradi."""
    client = _client()

    prompt = (
        f"Quyidagi fayl mazmuni va rasmlardagi barcha test savollarini HEMIS formatiga o'tkazib ber:\n\n{user_text.strip()}"
        if user_text.strip()
        else "Bu fayldagi barcha test savollarni HEMIS formatiga o'tkazib ber."
    )

    content: list = [{"type": "text", "text": prompt}]
    for img_bytes, mime in images[:10]:
        img_bytes, mime = _preprocess_image(img_bytes, mime)
        safe_mime = mime if mime in ('image/jpeg', 'image/png', 'image/gif', 'image/webp') else 'image/png'
        b64 = base64.standard_b64encode(img_bytes).decode()
        content.append({
            "type": "image_url",
            "image_url": {"url": f"data:{safe_mime};base64,{b64}"},
        })

    response = client.chat.completions.create(
        model=_model(),
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": content},
        ],
        temperature=0.1,
        max_tokens=8000,
    )
    return response.choices[0].message.content or ''


def _extract_text_from_binary(content: bytes, filename: str) -> str:
    """PDF/DOCX/XLSX fayllardan matn ajratib olish."""
    ext = ('.' + filename.rsplit('.', 1)[-1].lower()) if '.' in filename else ''

    if ext == '.pdf':
        try:
            import pdfplumber
            lines: list[str] = []
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                for page in pdf.pages:
                    try:
                        text = page.extract_text(layout=True, x_tolerance=3, y_tolerance=3) or ''
                    except TypeError:
                        text = page.extract_text(x_tolerance=3, y_tolerance=3) or ''
                    if text.strip():
                        lines.extend(text.splitlines())
                    else:
                        words = page.extract_words(x_tolerance=3, y_tolerance=3)
                        lines.extend(
                            w['text'] for w in sorted(words, key=lambda w: (w['top'], w['x0']))
                        )
            return '\n'.join(lines)
        except Exception as e:
            logger.warning("PDF text extraction failed: %s", e)

    elif ext == '.docx':
        try:
            from docx import Document
            doc = Document(io.BytesIO(content))
            lines = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        lines.append(' | '.join(cells))
            return '\n'.join(lines)
        except Exception as e:
            logger.warning("DOCX text extraction failed: %s", e)

    elif ext == '.xlsx':
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
            ws = wb.active
            lines = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c or '').strip() for c in row if c and str(c).strip()]
                if cells:
                    lines.append(' | '.join(cells))
            return '\n'.join(lines)
        except Exception as e:
            logger.warning("XLSX text extraction failed: %s", e)

    return content.decode('utf-8', errors='replace')


def run_for_file(file_obj, filename: str) -> list[ParsedQuestion]:
    """Fayl (rasm yoki matn/hujjat) bo'yicha GPT-4o orqali ParsedQuestion qaytaradi.

    DOCX: matn + ichki rasmlar (formula rasmlari) birgalikda Vision'ga yuboriladi.
    PDF:  avval matn ajratiladi; bo'sh bo'lsa sahifalar rasm sifatida yuboriladi.
    Rasm: to'g'ridan-to'g'ri Vision'ga.
    Boshqa: matn ajratilib text completion'ga.
    """
    ext = ('.' + filename.rsplit('.', 1)[-1].lower()) if '.' in filename else ''
    content = file_obj.read()

    if ext in IMAGE_EXTS:
        mime = MIME_MAP.get(ext, 'image/jpeg')
        raw = _call_image(content, mime)

    elif ext == '.docx':
        text, images = _extract_docx_content(content)
        if images:
            logger.info("DOCX ichida %d ta rasm topildi — multimodal so'rov yuborilmoqda", len(images))
            raw = _call_multimodal(text, images)
        else:
            raw = _call_text(
                f"Quyidagi matndan barcha savollarni HEMIS formatiga o'tkazib ber:\n\n{text}"
            )

    elif ext == '.pdf':
        text = _extract_text_from_binary(content, filename)
        if text.strip():
            raw = _call_text(
                f"Quyidagi matndan barcha savollarni HEMIS formatiga o'tkazib ber:\n\n{text}"
            )
        else:
            logger.info("PDF dan matn topilmadi — sahifalar rasm sifatida yuborilmoqda")
            page_images = _pdf_to_images(content)
            if page_images:
                raw = _call_multimodal('', page_images)
            else:
                raise RuntimeError(
                    "PDF dan matn o'qib bo'lmadi va sahifalarni rasm sifatida yuborib bo'lmadi. "
                    "pip install pymupdf buyrug'ini bajaring yoki faylni DOCX formatiga o'girib yuklang."
                )

    else:
        text = _extract_text_from_binary(content, filename)
        raw = _call_text(
            f"Quyidagi matndan barcha savollarni HEMIS formatiga o'tkazib ber:\n\n{text}"
        )

    return parse_hemis_text(raw)
